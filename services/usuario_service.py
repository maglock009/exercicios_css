from datetime import datetime
from io import BytesIO

import bcrypt
from openpyxl import Workbook

from repositories.auditoria_repository import AuditoriaRepository
from repositories.usuario_repository import UsuarioRepository


class UsuarioService:
    def __init__(
        self,
        usuario_repository: UsuarioRepository | None = None,
        auditoria_repository: AuditoriaRepository | None = None,
    ):
        self.usuarios = usuario_repository or UsuarioRepository()
        self.auditoria = auditoria_repository or AuditoriaRepository()

    @staticmethod
    def _hash_senha(senha: str) -> str:
        return bcrypt.hashpw(senha.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")

    @staticmethod
    def verificar_senha(senha: str, senha_hash: str) -> bool:
        return bcrypt.checkpw(senha.encode("utf-8"), senha_hash.encode("utf-8"))

    @staticmethod
    def formatar_data_br(valor) -> str:
        if valor is None or valor == "":
            return ""
        if isinstance(valor, str):
            texto = valor.strip()
            for formato in ("%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S"):
                try:
                    return datetime.strptime(texto[:19], formato).strftime("%d/%m/%Y")
                except ValueError:
                    continue
            if len(texto) == 10 and texto[2] == "/" and texto[5] == "/":
                return texto
            return texto
        if hasattr(valor, "strftime"):
            return valor.strftime("%d/%m/%Y")
        return str(valor)

    def autenticar(self, email: str, senha: str):
        usuario = self.usuarios.buscar_por_email(email)
        if not usuario:
            return None
        if not self.verificar_senha(senha, usuario["senha_hash"]):
            return None
        self.auditoria.registrar(
            usuario["id"],
            "LOGIN",
            "Acesso ao sistema",
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        )
        return usuario

    def listar(self):
        return self.usuarios.listar()

    def buscar_por_id(self, usuario_id: int):
        return self.usuarios.buscar_por_id(usuario_id)

    def cadastrar(self, email: str, senha: str, data_nascimento: str, papel: str, autor_id: int | None = None):
        existente = self.usuarios.buscar_por_email(email)
        if existente:
            raise ValueError("Já existe um usuário com este e-mail.")

        usuario_id = self.usuarios.criar(
            email=email.strip(),
            senha_hash=self._hash_senha(senha),
            data_nascimento=data_nascimento,
            papel=papel,
        )
        if autor_id:
            self.auditoria.registrar(
                autor_id,
                "CADASTRO",
                f"Usuário {email} criado",
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            )
        return usuario_id

    def atualizar(self, usuario_id: int, email: str, data_nascimento: str, papel: str, autor_id: int | None = None):
        atual = self.usuarios.buscar_por_id(usuario_id)
        if not atual:
            raise ValueError("Usuário não encontrado.")

        outro = self.usuarios.buscar_por_email(email.strip())
        if outro and outro["id"] != usuario_id:
            raise ValueError("Já existe um usuário com este e-mail.")

        self.usuarios.atualizar(usuario_id, email.strip(), data_nascimento, papel)
        if autor_id:
            self.auditoria.registrar(
                autor_id,
                "EDICAO",
                f"Usuário {email} atualizado",
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            )

    def atualizar_perfil(self, usuario_id: int, email: str, data_nascimento: str, nova_senha: str | None = None):
        atual = self.usuarios.buscar_por_id(usuario_id)
        if not atual:
            raise ValueError("Usuário não encontrado.")

        outro = self.usuarios.buscar_por_email(email.strip())
        if outro and outro["id"] != usuario_id:
            raise ValueError("Já existe um usuário com este e-mail.")

        self.usuarios.atualizar_perfil(usuario_id, email.strip(), data_nascimento)
        if nova_senha:
            self.usuarios.atualizar_senha(usuario_id, self._hash_senha(nova_senha))

        self.auditoria.registrar(
            usuario_id,
            "EDICAO",
            "Perfil atualizado",
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        )

    def excluir(self, usuario_id: int, autor_id: int | None = None):
        atual = self.usuarios.buscar_por_id(usuario_id)
        if not atual:
            raise ValueError("Usuário não encontrado.")
        if autor_id and autor_id == usuario_id:
            raise ValueError("Não é possível excluir o próprio usuário.")

        email = atual["email"]
        self.usuarios.excluir(usuario_id)
        if autor_id:
            self.auditoria.registrar(
                autor_id,
                "EXCLUSAO",
                f"Usuário {email} excluído",
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            )

    def gerar_relatorio_excel(self) -> BytesIO:
        workbook = Workbook()
        planilha = workbook.active
        planilha.title = "Usuarios"

        planilha["A1"] = "Relatórios de usuários"
        planilha["A3"] = "e-mail"
        planilha["B3"] = "data de nascimento"

        for indice, usuario in enumerate(self.usuarios.listar(), start=4):
            planilha.cell(row=indice, column=1, value=usuario["email"])
            planilha.cell(
                row=indice,
                column=2,
                value=self.formatar_data_br(usuario["data_nascimento"]),
            )

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer
